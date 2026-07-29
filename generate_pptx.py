from pptx import Presentation
from pptx.util import Inches, Pt

def add_slide(prs, title, headline, points):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    title_shape = slide.shapes.title
    title_shape.text = title
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    # Add headline
    p = tf.text = headline
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(24)
    
    # Add points
    for point in points:
        p = tf.add_paragraph()
        p.text = point
        p.level = 1
        p.font.size = Pt(18)

prs = Presentation()

# Slide 1
title1 = "Today's Accomplishments"
headline1 = "Taking pydpvz from a basic I/O wrapper to a full-fledged distributed processing ecosystem."
points1 = [
    "New Tools Built: Successfully developed and integrated dpvtkextract, dpvtksplice, dpvtkfilter (in-memory parallel algorithms), and dpvtkvideo (FFmpeg-powered MPI animation).",
    "Automated CI/CD: Completed Phase 17 by implementing a rigorous pytest integration suite that tests parallel end-to-end reading, filtering, and .dpvtk generation.",
    "Open-Source Ready: Fully sanitized the repository (w/ strict .gitignore), wrote a professional README.md, generalized the MPI setup scripts, and published to GitHub."
]
add_slide(prs, title1, headline1, points1)

# Slide 2
title2 = "Issues & Technical Hurdles Overcome"
headline2 = "Debugging distributed systems and ParaView pipelines with AI."
points2 = [
    "The 'N vs M' Processor Trap: Identified a critical bug where reading a 40-rank dataset on a 16-rank batch job silently dropped chunks 16-39. We resolved this by implementing a dynamic round-robin block distribution loop.",
    "Pipeline Type Crashes: Complex ParaView filters were returning raw vtkDataObjects instead of standard datasets, crashing the in-memory serializer. We fixed this by deeply interrogating the proxy's GetClientSideObject().",
    "MPI ABI Deadlocks: Prevented catastrophic deadlocks for future users by embedding a dynamic ldd MPI-detector into the setup scripts."
]
add_slide(prs, title2, headline2, points2)

# Slide 3
title3 = "AI Tooling & Resource Footprint"
headline3 = "Leveraging Agentic AI to rapidly prototype complex architectures."
points3 = [
    "Token Usage: ~240,000 tokens today (July 28) vs ~500,000 tokens yesterday.",
    "Reduced token usage today by shifting from 'high-effort' heavy reasoning models to 'low-effort' lightning-fast models for mechanical, repetitive tasks.",
    "Successfully created an AGENTS.md file to bootstrap future AI sessions instantly, ensuring new agents don't hallucinate or fall into known MPI traps."
]
add_slide(prs, title3, headline3, points3)

prs.save('/workspaces/AllVibesDemo/workshop_update.pptx')
print("Successfully generated workshop_update.pptx")
